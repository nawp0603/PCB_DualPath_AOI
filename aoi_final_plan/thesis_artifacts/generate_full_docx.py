# -*- coding: utf-8 -*-
"""
Main Publisher Script: Generates the complete publication-grade Word document (.docx)
for the Dual-Path Industrial AOI Thesis & Portable Release Package.

Outputs:
1. D:\FPTU\KLTN\PCB_DualPath_AOI\aoi_final_plan\thesis_artifacts\PCB_DualPath_AOI_Full_Documentation.docx
2. D:\FPTU\KLTN\PCB_DualPath_AOI\PCB_DualPath_AOI_Documentation.docx
"""
import shutil
import sys
from pathlib import Path
from docx import Document

THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

from docx_style_utils import configure_document
from docx_chapters_1_3 import (
    build_cover_and_executive_summary,
    build_chapter_1,
    build_chapter_2,
    build_chapter_3,
)
from docx_chapters_4_6 import (
    build_chapter_4,
    build_chapter_5,
    build_chapter_6,
)


def main() -> None:
    fig_dir = THIS_DIR / "figures"
    out_docx_artifacts = THIS_DIR / "PCB_DualPath_AOI_Full_Documentation.docx"
    project_root = THIS_DIR.parent.parent
    out_docx_root = project_root / "PCB_DualPath_AOI_Documentation.docx"

    doc = Document()
    configure_document(doc)

    # Build Cover Page + Executive Summary + Chapters 1..6
    build_cover_and_executive_summary(doc, fig_dir)
    doc.add_page_break()
    build_chapter_1(doc, fig_dir)
    doc.add_page_break()
    build_chapter_2(doc, fig_dir)
    doc.add_page_break()
    build_chapter_3(doc, fig_dir)
    doc.add_page_break()
    build_chapter_4(doc, fig_dir)
    doc.add_page_break()
    build_chapter_5(doc, fig_dir)
    doc.add_page_break()
    build_chapter_6(doc, fig_dir)

    doc.save(str(out_docx_artifacts))
    shutil.copy2(str(out_docx_artifacts), str(out_docx_root))

    # Post-Generation Verification
    verify_doc = Document(str(out_docx_artifacts))
    n_paragraphs = len(verify_doc.paragraphs)
    h1_count = sum(1 for p in verify_doc.paragraphs if p.style.name == "Heading 1")
    h2_count = sum(1 for p in verify_doc.paragraphs if p.style.name == "Heading 2")
    h3_count = sum(1 for p in verify_doc.paragraphs if p.style.name == "Heading 3")
    n_tables = len(verify_doc.tables)
    n_shapes = len(verify_doc.inline_shapes)

    size_artifacts_kb = out_docx_artifacts.stat().st_size / 1024.0
    size_root_kb = out_docx_root.stat().st_size / 1024.0

    print("=" * 72)
    print("DOCX GENERATION & VERIFICATION COMPLETE")
    print("=" * 72)
    print(f"1. Artifact DOCX : {out_docx_artifacts} ({size_artifacts_kb:.1f} KB)")
    print(f"2. Root Copy DOCX: {out_docx_root} ({size_root_kb:.1f} KB)")
    print(f"3. Paragraphs    : {n_paragraphs}")
    print(f"4. Headings      : Total={h1_count + h2_count + h3_count} (H1={h1_count}, H2={h2_count}, H3={h3_count})")
    print(f"5. Styled Tables : {n_tables}")
    print(f"6. Inline Figures: {n_shapes}")
    print("=" * 72)


if __name__ == "__main__":
    main()
