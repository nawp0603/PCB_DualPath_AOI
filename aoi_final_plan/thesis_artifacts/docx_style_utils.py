# -*- coding: utf-8 -*-
"""
Publication-Grade Styling Utilities for python-docx
Generates A4 academic layout, Navy/Slate headings, Zebra-striped tables with
highlighted proposed rows, Callout boxes (#F4F6F9 + thick left border), and
centered high-DPI figure embedding.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

NAVY_HEX = "1B365D"
SLATE_HEX = "2C5282"
SUB_HEX = "1E3A8A"
CALLOUT_BG_HEX = "F4F6F9"
ZEBRA_HEX = "F8FAFC"
HIGHLIGHT_HEX = "E0F2FE"
BORDER_HEX = "CBD5E1"

NAVY_RGB = RGBColor(0x1B, 0x36, 0x5D)
SLATE_RGB = RGBColor(0x2C, 0x52, 0x82)
SUB_RGB = RGBColor(0x1E, 0x3A, 0x8A)
BODY_RGB = RGBColor(0x1E, 0x29, 0x3B)
MUTED_RGB = RGBColor(0x47, 0x55, 0x69)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)

FONT_MAIN = "Times New Roman"
FONT_CODE = "Consolas"


def set_cell_shading(cell, color_hex: str) -> None:
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color_hex: str = BORDER_HEX) -> None:
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{NAVY_HEX}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{NAVY_HEX}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def configure_document(doc: Document) -> None:
    """Configure A4 page margins, default fonts, and header/footer."""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.2)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(
            "KLTN: Hệ thống Kiểm tra Ngoại quan Công nghiệp Hai Nhánh (Dual-Path Industrial AOI) — Mạch PCB Trần"
        )
        hrun.font.name = FONT_MAIN
        hrun.font.size = Pt(9)
        hrun.font.italic = True
        hrun.font.color.rgb = MUTED_RGB

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(
            "FPT University — Tài liệu Kỹ thuật & Báo cáo Nghiệm thu Toàn diện (100% Số liệu Thực nghiệm Đã Xác minh)"
        )
        frun.font.name = FONT_MAIN
        frun.font.size = Pt(9)
        frun.font.color.rgb = MUTED_RGB

    normal_style = doc.styles["Normal"]
    normal_style.font.name = FONT_MAIN
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = BODY_RGB
    normal_style.paragraph_format.line_spacing = 1.22
    normal_style.paragraph_format.space_after = Pt(6)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_heading(level=1)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_MAIN
    run.font.size = Pt(15.5)
    run.font.bold = True
    run.font.color.rgb = NAVY_RGB

    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="12" w:space="4" w:color="{NAVY_HEX}"/>'
        f"</w:pBdr>"
    )
    pPr.append(pBdr)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_heading(level=2)
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_MAIN
    run.font.size = Pt(13.0)
    run.font.bold = True
    run.font.color.rgb = SLATE_RGB


def add_h3(doc: Document, text: str) -> None:
    p = doc.add_heading(level=3)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_MAIN
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = SUB_RGB


def add_p(doc: Document, text: str, bold_prefix: str = "", italic: bool = False, justify: bool = True) -> None:
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = FONT_MAIN
        r_bold.font.size = Pt(11.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = BODY_RGB
    r = p.add_run(text)
    r.font.name = FONT_MAIN
    r.font.size = Pt(11.5)
    r.font.italic = italic
    r.font.color.rgb = BODY_RGB


def add_bullet(doc: Document, bold_prefix: str, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = FONT_MAIN
        rb.font.size = Pt(11.5)
        rb.font.bold = True
        rb.font.color.rgb = NAVY_RGB
    r = p.add_run(text)
    r.font.name = FONT_MAIN
    r.font.size = Pt(11.5)
    r.font.color.rgb = BODY_RGB


def add_callout(doc: Document, title: str, lines: list[str], is_code: bool = False) -> None:
    """Create a shaded callout box (#F4F6F9) with thick Navy left border."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.18

    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CALLOUT_BG_HEX}"/>')
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="10" w:color="{NAVY_HEX}"/>'
        f'<w:top w:val="single" w:sz="4" w:space="4" w:color="{BORDER_HEX}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="4" w:color="{BORDER_HEX}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="6" w:color="{BORDER_HEX}"/>'
        f"</w:pBdr>"
    )
    pPr.append(shd)
    pPr.append(pBdr)

    if title:
        rt = p.add_run(title + "\n")
        rt.font.name = FONT_MAIN
        rt.font.size = Pt(11.0)
        rt.font.bold = True
        rt.font.color.rgb = NAVY_RGB

    for idx, line in enumerate(lines):
        suffix = "\n" if idx < len(lines) - 1 else ""
        rl = p.add_run(line + suffix)
        rl.font.name = FONT_CODE if is_code else FONT_MAIN
        rl.font.size = Pt(10.0 if is_code else 11.0)
        rl.font.color.rgb = BODY_RGB


def add_styled_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    highlight_keywords: tuple[str, ...] = ("Dual-Path", "System 3", "Hệ thống 3", "ĐẠT", "Trung bình Toàn cục"),
) -> None:
    """Create a publication-grade table with Navy header, zebra striping, and highlighted rows."""
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.keep_with_next = True
        cr = cp.add_run(caption)
        cr.font.name = FONT_MAIN
        cr.font.size = Pt(11.0)
        cr.font.bold = True
        cr.font.color.rgb = NAVY_RGB

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header row
    hdr_cells = table.rows[0].cells
    for j, htext in enumerate(headers):
        cell = hdr_cells[j]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY_HEX)
        set_cell_margins(cell, top=110, bottom=110, left=110, right=110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(htext)
        run.font.name = FONT_MAIN
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = WHITE_RGB

    # Data rows
    for i, row_data in enumerate(rows):
        row_cells = table.rows[i + 1].cells
        row_joined = " ".join(str(x) for x in row_data)
        is_highlight = any(kw in row_joined for kw in highlight_keywords)
        bg_hex = HIGHLIGHT_HEX if is_highlight else (ZEBRA_HEX if i % 2 == 1 else "FFFFFF")

        for j, val in enumerate(row_data):
            cell = row_cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, bg_hex)
            set_cell_margins(cell, top=85, bottom=85, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            run = p.add_run(str(val))
            run.font.name = FONT_MAIN
            run.font.size = Pt(9.5)
            run.font.bold = is_highlight or (j == 0)
            run.font.color.rgb = NAVY_RGB if is_highlight else BODY_RGB

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def add_figure(doc: Document, img_path: Path, caption: str, width_inches: float = 5.8) -> None:
    """Embed a high-DPI PNG figure centered with academic caption below."""
    if not img_path.exists():
        add_callout(doc, "Cảnh báo thiếu tệp hình ảnh", [f"Không tìm thấy: {img_path}"])
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(str(img_path), width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = FONT_MAIN
    r_cap.font.size = Pt(10.0)
    r_cap.font.italic = True
    r_cap.font.bold = True
    r_cap.font.color.rgb = SLATE_RGB
